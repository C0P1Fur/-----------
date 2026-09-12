# -*- coding: utf-8 -*-
"""微网与外部电网电力调控策略——公共函数"""
from __future__ import annotations

import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from scipy.optimize import milp, LinearConstraint, Bounds
from scipy.sparse import lil_matrix

# 功率统一使用kW且电量统一使用kWh，因此十分钟功率乘DT才得到时段电量
DT = 1.0 / 6.0
ETA_C = 0.9
ETA_D = 0.9
# E表示电池内部储电量，容量边界与充放电量的微网侧口径不同
E_MIN = 1200.0
E_MAX = 10800.0
P_MAX_E = 5000.0 * DT  # 每10分钟最大充/放电量 kWh
# 岭惩罚固定为10，用于减轻滞后特征相关性造成的系数不稳定
RIDGE_LAMBDA = 10.0
N_SLOT = 144
DAYS = 365
# 日期索引从2025年1月1日的0开始，2月1日对应31
BASE_DATE = date(2025, 1, 1)
RELEASE_HOURS = np.array([0, 6, 12, 18], dtype=int)


# ============================================================================
# Excel 数据读取
# ============================================================================
# 读取附件1 返回price, load, pv
def read_attachment1(path: str):
    print(f"[读取] 正在读取附件1：{path}", flush=True)

    # 只读模式减少内存占用，data_only读取已有公式缓存而不会重新计算Excel公式
    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False
    )

    ws = wb["Sheet1"]

    data = []
    for row in ws.iter_rows(
        min_row=2,
        max_row=145,
        min_col=2,
        max_col=4,
        values_only=True
    ):
        data.append(row)

    wb.close()

    # 将附件1的三列依次解释为电价与负荷及光伏，并转换为浮点数组
    arr = np.asarray(data, dtype=float)

    price = arr[:, 0]
    load = arr[:, 1]
    pv = arr[:, 2]

    print("[读取] 附件1读取完成", flush=True)

    return price, load, pv


# 读取附件2 返回load(act), pv(act)
def read_attachment2(path: str):

    print(f"[读取] 正在打开附件2：{path}", flush=True)

    # 只读模式减少内存占用，data_only读取已有公式缓存而不会重新计算Excel公式
    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False
    )

    print("[读取] Excel 已打开，正在读取小区负载...", flush=True)

    ws_load = wb["小区负载"]

    load_data = []

    for row in ws_load.iter_rows(
        min_row=2,
        max_row=366,
        min_col=2,
        max_col=145,
        values_only=True
    ):
        load_data.append(row)

    print("[读取] 小区负载读取完成，正在读取光伏...", flush=True)

    ws_pv = wb["光伏发电实际功率"]

    pv_data = []

    for row in ws_pv.iter_rows(
        min_row=2,
        max_row=366,
        min_col=2,
        max_col=145,
        values_only=True
    ):
        pv_data.append(row)

    wb.close()

    print("[读取] Excel 文件已关闭，正在转换 NumPy 数组...", flush=True)

    # 每行对应一天且每列对应十分钟时段，全年数据形状为365乘144
    load = np.asarray(load_data, dtype=float)
    pv = np.asarray(pv_data, dtype=float)

    print(
        f"[读取] 附件2读取完成："
        f"load.shape={load.shape}, "
        f"pv.shape={pv.shape}",
        flush=True
    )

    return load, pv


"""
读取附件3
返回：
    forecast: ndarray, shape = (365, 4, 24)

    forecast[d, k, h]
    d = 0~364，对应 2025-01-01 ~ 2025-12-31
    k = 0,1,2,3，对应 0:00、6:00、12:00、18:00
    h = 0~23，对应未来1~24小时
"""
def read_attachment3(path: str):

    print(f"[读取] 正在读取附件3：{path}", flush=True)

    # 只读模式减少内存占用，data_only读取已有公式缓存而不会重新计算Excel公式
    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False
    )

    ws = wb["Sheet1"]

    # 365天 × 每天4次发布 × 未来24小时
    forecast = np.zeros((365, 4, 24), dtype=float)

    # 当前读取方式依赖预报行按每天四次发布排列，日期和版本由行号推算
    row_index = 0

    for row in ws.iter_rows(
        min_row=2,
        max_row=1461,
        min_col=3,
        max_col=26,
        values_only=True
    ):
        day = row_index // 4
        release = row_index % 4

        # None 安全处理
        # 按照下面的分支处理空值，再将当前行转换为数值列表
        values = [
            0.0 if v is None else float(v)
            for v in row
        ]

        forecast[day, release, :] = values

        row_index += 1

    wb.close()

    # 此处检查循环行数，固定读取范围中的空白行也会参与计数
    if row_index != 1460:
        raise ValueError(
            f"附件3读取行数异常：实际 {row_index} 行，预期 1460 行"
        )

    print(
        f"[读取] 附件3完成，forecast.shape={forecast.shape}",
        flush=True
    )

    return forecast

"""
读取附件4
返回：
    price[d, t]
    d = 0~364
    t = 0~143
"""
def read_attachment4(path: str):

    print(f"[读取] 正在读取附件4：{path}", flush=True)

    # 只读模式减少内存占用，data_only读取已有公式缓存而不会重新计算Excel公式
    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False
    )

    ws = wb["Sheet1"]

    price_data = []

    for row in ws.iter_rows(
        min_row=2,
        max_row=366,
        min_col=2,
        max_col=145,
        values_only=True
    ):
        # 按照下面的分支处理空值，再将当前行转换为数值列表
        values = [
            np.nan if v is None else float(v)
            for v in row
        ]

        price_data.append(values)

    wb.close()

    price = np.asarray(price_data, dtype=float)

    if price.shape != (365, 144):
        raise ValueError(
            f"附件4维度异常：实际 {price.shape}，预期 (365, 144)"
        )

    # 空白电价已转换为NaN，此处发现后直接停止读取
    if np.isnan(price).any():
        raise ValueError("附件4中存在空白电价数据")

    print(
        f"[读取] 附件4完成，price.shape={price.shape}",
        flush=True
    )

    return price

# ============================================================================
# MILP 优化与真实执行
# ============================================================================
"""决策变量顺序[q,c,d,E,y] 返回计划购电量"""
def solve_nominal(net_pred_kw: np.ndarray, reserve_kw: np.ndarray,
                  price: np.ndarray, e_start: float, e_terminal: float = 6000.0):
    n = len(net_pred_kw) # 获得单决策变量数量
    nv = 5*n + 1 # 总变量的数量
    # 变量按q与c及d和E及y分块排列，其中E含n加1个边界状态
    iq, ic, idd, iE, iy = 0, n, 2*n, 3*n, 4*n + 1 # 购电量 充电量 放电量 储电量 状态变量在数组中的位置

    # 输出
    # 目标系数初始为零，仅对随后显式赋价的变量计费
    obj = np.zeros(nv)
    obj[:n] = price
    # lb和ub规定变量边界，购电与充放电默认非负而储能边界另行覆盖
    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)
    ub[ic:ic+n] = P_MAX_E
    ub[idd:idd+n] = P_MAX_E
    lb[iE:iE+n+1] = E_MIN
    ub[iE:iE+n+1] = E_MAX
    ub[iy:iy+n] = 1.0
    # 0表示连续变量，模式变量设为整数并结合零到一边界形成二元变量
    integrality = np.zeros(nv)
    integrality[iy:iy+n] = 1

    # 日前约束共4n加2条，由状态递推与供需及两类互斥和首末状态组成
    A = lil_matrix((4*n + 2, nv))
    # 每行约束统一写成lo不大于Ax不大于hi，令上下界相同即可表示等式
    lo = np.full(4*n + 2, -np.inf)
    hi = np.full(4*n + 2, np.inf)
    r = 0
    for t in range(n):
        # 本行状态方程为E_next减E减ETA_C乘c加d除ETA_D等于零
        A[r, iE+t+1] = 1
        A[r, iE+t] = -1
        A[r, ic+t] = -ETA_C
        A[r, idd+t] = 1 / ETA_D
        lo[r] = hi[r] = 0
        r += 1
    # 把当前真实储电量固定为规划初值，r随后移到下一条约束
    A[r, iE] = 1; lo[r] = hi[r] = e_start; r += 1
    # 末端约束作用于名义预测轨迹，真实运行后的状态可能与该目标不同
    A[r, iE+n] = 1; lo[r] = hi[r] = e_terminal; r += 1

    # 将预测净负荷加备用功率转换为每时段需要保障的名义电量
    req = (net_pred_kw + reserve_kw) * DT
    for t in range(n):
        # 供需约束为q减c加d不小于req，允许存在未利用的剩余供给
        A[r, t] = 1
        A[r, ic+t] = -1
        A[r, idd+t] = 1
        lo[r] = req[t]
        r += 1
    for t in range(n):
        # c不超过M乘y且d不超过M乘一减y，因此同一时段不能同时充放电
        A[r, ic+t] = 1; A[r, iy+t] = -P_MAX_E; hi[r] = 0; r += 1
        A[r, idd+t] = 1; A[r, iy+t] = P_MAX_E; hi[r] = P_MAX_E; r += 1

    # 将稀疏约束与变量类型交给MILP求解器，15秒时限与1e-7相对间隙控制求解停止
    res = milp(
        obj,
        integrality=integrality,
        bounds=Bounds(lb, ub),
        constraints=LinearConstraint(A.tocsr(), lo, hi),
        options={"time_limit": 15, "mip_rel_gap": 1e-7},
    )
    # 未成功时主动报错，避免后续使用不满足当前成功条件的解
    if not res.success:
        raise RuntimeError(f"MILP求解失败: {res.message}")
    x = res.x
    # 按变量块返回购电与充放电及储能轨迹和名义费用，状态数组比时段数多一个元素
    return x[:n], x[ic:ic+n], x[idd:idd+n], x[iE:iE+n+1], float(res.fun)

"""考虑计划购电量和实际购电量的MILP"""
def solve_adjust(q0: np.ndarray, net_pred_kw: np.ndarray, reserve_kw: np.ndarray,
                 price: np.ndarray, e_start: float):
    n = len(q0)
    # 调整模型增加增购u与减购v两个变量块，共7n加1个变量
    iq, iu, iv, ic, idd, iE, iy = 0, n, 2*n, 3*n, 4*n, 5*n, 6*n + 1
    nv = 7*n + 1
    # 目标系数初始为零，仅对随后显式赋价的变量计费
    obj = np.zeros(nv)
    # 增购按1.5倍价格计费，减购负系数对应全额退费再收50%违约费的假设
    obj[iu:iu+n] = 1.5 * price
    obj[iv:iv+n] = -0.5 * price
    # lb和ub规定变量边界，购电与充放电默认非负而储能边界另行覆盖
    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)
    # 减购量不能超过原日前承诺量，q0始终是原计划而非上一次更新版本
    ub[iv:iv+n] = q0
    ub[ic:ic+n] = P_MAX_E
    ub[idd:idd+n] = P_MAX_E
    lb[iE:iE+n+1] = E_MIN
    ub[iE:iE+n+1] = E_MAX
    ub[iy:iy+n] = 1.0
    # 0表示连续变量，模式变量设为整数并结合零到一边界形成二元变量
    integrality = np.zeros(nv)
    integrality[iy:iy+n] = 1

    # 调整模型比日前模型多n条计划关系约束，因此共5n加2条约束
    A = lil_matrix((5*n + 2, nv))
    lo = np.full(5*n + 2, -np.inf)
    hi = np.full(5*n + 2, np.inf)
    r = 0
    for t in range(n):
        # 用q减u加v等于q0连接最终购电量和相对原计划的增减量
        A[r, iq+t] = 1; A[r, iu+t] = -1; A[r, iv+t] = 1
        lo[r] = hi[r] = q0[t]
        r += 1
    for t in range(n):
        # 更新模型沿用充电乘效率与放电除效率的内部储能递推
        A[r, iE+t+1] = 1; A[r, iE+t] = -1
        A[r, ic+t] = -ETA_C; A[r, idd+t] = 1/ETA_D
        lo[r] = hi[r] = 0
        r += 1
    # 把当前真实储电量固定为规划初值，r随后移到下一条约束
    A[r, iE] = 1; lo[r] = hi[r] = e_start; r += 1
    # 本函数将剩余规划的日末名义储电量固定为6000千瓦时
    A[r, iE+n] = 1; lo[r] = hi[r] = 6000.0; r += 1
    # 将预测净负荷加备用功率转换为每时段需要保障的名义电量
    req = (net_pred_kw + reserve_kw) * DT
    for t in range(n):
        # 用调整后的常规购电量保障剩余预测净负荷和备用裕量
        A[r, iq+t] = 1; A[r, ic+t] = -1; A[r, idd+t] = 1
        lo[r] = req[t]
        r += 1
    for t in range(n):
        # c不超过M乘y且d不超过M乘一减y，因此同一时段不能同时充放电
        A[r, ic+t] = 1; A[r, iy+t] = -P_MAX_E; hi[r] = 0; r += 1
        A[r, idd+t] = 1; A[r, iy+t] = P_MAX_E; hi[r] = P_MAX_E; r += 1

    # 将稀疏约束与变量类型交给MILP求解器，15秒时限与1e-7相对间隙控制求解停止
    res = milp(
        obj,
        integrality=integrality,
        bounds=Bounds(lb, ub),
        constraints=LinearConstraint(A.tocsr(), lo, hi),
        options={"time_limit": 15, "mip_rel_gap": 1e-7},
    )
    # 未成功时主动报错，避免后续使用不满足当前成功条件的解
    if not res.success:
        raise RuntimeError(f"调整MILP求解失败: {res.message}")
    # 只返回调整后的常规购电计划，实际充放电由执行函数根据真实供需重新决定
    return res.x[:n]


def execute_segment(q: np.ndarray, load: np.ndarray, pv: np.ndarray, e_start: float):
    """真实执行：盈余先充电，缺口先放电，仍不足则紧急购电。"""
    n = len(q)
    ch = np.zeros(n); dis = np.zeros(n); emergency = np.zeros(n); waste = np.zeros(n)
    E = np.zeros(n + 1); E[0] = e_start
    for t in range(n):
        # 正余额表示供给盈余可充电，负余额表示需要放电或紧急补购
        balance = q[t] + (pv[t] - load[t]) * DT
        if balance >= 0:
            # 充电量同时受盈余电量与充电功率及电池剩余容量限制
            ch[t] = max(0.0, min(balance, P_MAX_E, (E_MAX - E[t]) / ETA_C))
            waste[t] = max(0.0, balance - ch[t])
        else:
            # 放电量同时受缺口与功率及剩余可用储能限制，仍不足的部分才紧急补购
            dis[t] = max(0.0, min(-balance, P_MAX_E, (E[t] - E_MIN) * ETA_D))
            emergency[t] = max(0.0, -balance - dis[t])
        # 按微网侧充放电量更新电池内部状态，损耗只在状态方程中折算一次
        E[t+1] = E[t] + ETA_C * ch[t] - dis[t] / ETA_D
    # 返回实际执行量和n加1个储能边界，waste可能包含已购但未使用的电量
    return ch, dis, emergency, waste, E

# ============================================================================
# 预测模型
# ============================================================================
"""滚动岭回归 12维特征 60天窗口。"""
def build_load_pv_predictions(load: np.ndarray, pv: np.ndarray,
                              init_load: np.ndarray, init_pv: np.ndarray):
    D, N = load.shape
    feat = np.zeros((D, N, 12), dtype=float)
    t = np.arange(N)
    theta = 2 * np.pi * t / N
    # 七天历史充分后才构造滞后特征，所有观测切片均截止当前日期之前
    for d in range(7, D):
        # 前六维使用前一天与前七天及近七天均值的负荷和光伏信息
        feat[d, :, 0] = load[d - 1]
        feat[d, :, 1] = pv[d - 1]
        feat[d, :, 2] = load[d - 7]
        feat[d, :, 3] = pv[d - 7]
        feat[d, :, 4] = load[d - 7:d].mean(axis=0)
        feat[d, :, 5] = pv[d - 7:d].mean(axis=0)
        # 后六维使用日内一阶与二阶周期及七天周期的正弦余弦特征
        feat[d, :, 6] = np.sin(theta)
        feat[d, :, 7] = np.cos(theta)
        feat[d, :, 8] = np.sin(2 * theta)
        feat[d, :, 9] = np.cos(2 * theta)
        feat[d, :, 10] = math.sin(2 * np.pi * (d % 7) / 7)
        feat[d, :, 11] = math.cos(2 * np.pi * (d % 7) / 7)

    pred_l = np.zeros_like(load)
    pred_g = np.zeros_like(pv)
    # 首日没有历史记录，因此使用附件1的典型日负荷与光伏启动预测
    pred_l[0] = init_load
    pred_g[0] = init_pv

    # 前15天采用此前最多七天同一时段的均值，随后才启用岭回归
    for d in range(1, 15):
        s = max(0, d - 7)
        pred_l[d] = load[s:d].mean(axis=0)
        pred_g[d] = pv[s:d].mean(axis=0)

    for d in range(15, D):
        # 训练窗口最多60天且跳过无法构造七天滞后的初始日期
        s = max(7, d - 60)
        # 把历史日期和时段合并为样本维度，每个训练样本保留12个特征
        X = feat[s:d].reshape(-1, 12)
        Y = np.column_stack([load[s:d].reshape(-1), pv[s:d].reshape(-1)])
        # 均值与标准差只从本次历史训练集估计，避免标准化使用未来信息
        mu = X.mean(axis=0)
        sd = X.std(axis=0)
        # 常量特征用尺度1避免除零，不改变其中心化后为零的结果
        sd[sd == 0] = 1.0
        Z = (X - mu) / sd
        ym = Y.mean(axis=0)
        # 求解带岭惩罚的正规方程并中心化响应，从而保留不受惩罚的截距
        beta = np.linalg.solve(
            Z.T @ Z + RIDGE_LAMBDA * np.eye(12),
            Z.T @ (Y - ym),
        )
        # 使用同一训练尺度转换当天特征，再乘回归系数并加回响应均值
        P = ym + ((feat[d] - mu) / sd) @ beta
        pred_l[d] = np.maximum(P[:, 0], 0.0)
        pred_g[d] = np.maximum(P[:, 1], 0.0)
        # 若此前七天某时段光伏都为零则预测归零，不根据未来数据判断夜间
        pred_g[d][np.all(pv[d - 7:d] == 0, axis=0)] = 0.0
    return pred_l, pred_g

"""滚动岭回归电价预测 9维"""
def build_price_predictions(price: np.ndarray, init_price: np.ndarray):
    D, N = price.shape
    feat = np.zeros((D, N, 9), dtype=float)
    t = np.arange(N)
    theta = 2 * np.pi * t / N
    # 七天历史充分后才构造滞后特征，所有观测切片均截止当前日期之前
    for d in range(7, D):
        # 电价前三维使用前一天与前七天及近七天均值，后六维表达周期变化
        feat[d, :, 0] = price[d - 1]
        feat[d, :, 1] = price[d - 7]
        feat[d, :, 2] = price[d - 7:d].mean(axis=0)
        feat[d, :, 3] = np.sin(theta)
        feat[d, :, 4] = np.cos(theta)
        feat[d, :, 5] = np.sin(2 * theta)
        feat[d, :, 6] = np.cos(2 * theta)
        feat[d, :, 7] = math.sin(2 * np.pi * (d % 7) / 7)
        feat[d, :, 8] = math.cos(2 * np.pi * (d % 7) / 7)

    pred = np.zeros_like(price)
    pred[0] = init_price
    # 前15天采用此前最多七天同一时段的均值，随后才启用岭回归
    for d in range(1, 15):
        s = max(0, d - 7)
        pred[d] = price[s:d].mean(axis=0)
    for d in range(15, D):
        # 训练窗口最多60天且跳过无法构造七天滞后的初始日期
        s = max(7, d - 60)
        # 将日期和时段展平成训练样本，每个电价样本具有9维特征
        X = feat[s:d].reshape(-1, 9)
        y = price[s:d].reshape(-1)
        # 均值与标准差只从本次历史训练集估计，避免标准化使用未来信息
        mu = X.mean(axis=0)
        sd = X.std(axis=0)
        # 常量特征用尺度1避免除零，不改变其中心化后为零的结果
        sd[sd == 0] = 1.0
        Z = (X - mu) / sd
        ym = y.mean()
        # 求解带岭惩罚的正规方程并中心化响应，从而保留不受惩罚的截距
        beta = np.linalg.solve(
            Z.T @ Z + RIDGE_LAMBDA * np.eye(9), Z.T @ (y - ym)
        )
        # 预测价格最低设为0.001元每千瓦时，以保留当前优化模型的正价格条件
        pred[d] = np.maximum(ym + ((feat[d] - mu) / sd) @ beta, 0.001)
    return pred

'''裕量计算'''
def reserve_from_residuals(residuals: np.ndarray, day: int, quantile: float = 0.5):
    if day == 0:
        return np.zeros(residuals.shape[1], dtype=float)
    # 只用过去最多28天同一时段误差，正误差表示实际净负荷高于预测
    hist = residuals[max(0, day - 28):day]
    # 逐时段计算经验分位数并把负值截为零，裕量单位仍为kW
    return np.maximum(np.quantile(hist, quantile, axis=0), 0.0)

# ============================================================================
# 问题3/4的日内光伏预测与反馈
# ============================================================================
def build_corrected_pv_forecasts(F: np.ndarray, pv_actual: np.ndarray):
    """附件3：同版本/同提前量28天历史偏差 + n/(n+7)收缩校正。"""
    D = F.shape[0]
    ferr = np.full_like(F, np.nan, dtype=float)

    def actual_at_hour(day: int, hour_abs: int):
        # 将发布日期和可能跨日的目标小时转成绝对时间编号
        total = day * 24 + hour_abs
        # 将24时终点归入前一天末时段，与实测十分钟终点索引对齐
        target_day = (total - 1) // 24
        local_hour = total - target_day * 24
        idx = int(local_hour * 6 - 1)
        if target_day < 0 or target_day >= D:
            return np.nan
        return pv_actual[target_day, idx]

    for d in range(D):
        for k, rh in enumerate(RELEASE_HOURS):
            for h in range(1, 25):
                a = actual_at_hour(d, int(rh + h))
                if np.isfinite(a):
                    # 误差定义为预测减实际，因此正均值表示预报系统性偏高
                    ferr[d, k, h-1] = F[d, k, h-1] - a

    corrected = np.zeros_like(F)
    for d in range(D):
        for k in range(4):
            for h in range(24):
                # 只抽取此前28天相同发布版本和相同提前量的有效误差
                e = ferr[max(0, d-28):d, k, h]
                e = e[np.isfinite(e)]
                n = len(e)
                bias = float(e.mean()) if n else 0.0
                # 历史样本越少修正幅度越小，n为零时不校正预报
                shrink = n / (n + 7.0) if n else 0.0
                # 从当前预报减去收缩后的历史偏差并保证光伏功率非负
                corrected[d, k, h] = max(F[d, k, h] - shrink * bias, 0.0)
    return corrected


def pv_forecast_10min(corrected: np.ndarray, pv_actual: np.ndarray, day: int, version: int):
    s = int(RELEASE_HOURS[version] * 6)
    # 插值锚点使用发布前最近实测光伏，只有首日零点无历史时设为零
    if day == 0 and s == 0:
        anchor = 0.0
    elif s == 0:
        anchor = pv_actual[day-1, 143]
    else:
        anchor = pv_actual[day, s-1]
    # 将已知锚点与同一版本的24个小时预报拼接为插值节点
    knots = np.r_[anchor, corrected[day, version]]
    out = np.zeros(N_SLOT - s)
    for j, t in enumerate(range(s, N_SLOT)):
        # 按区间终点计算相对发布时刻的分钟数，第一个目标位于发布后十分钟
        minutes = (t - s + 1) * 10
        h0 = minutes // 60
        if minutes % 60 == 0:
            out[j] = knots[h0]
        else:
            f = (minutes % 60) / 60.0
            # 按前后小时节点的时间距离线性加权得到十分钟预测功率
            out[j] = (1-f) * knots[h0] + f * knots[h0+1]
    return out


def build_version_residuals(pred_load: np.ndarray, load: np.ndarray, pv: np.ndarray,
                            corrected: np.ndarray):
    out = np.full((4, DAYS, N_SLOT), np.nan)
    for k, rh in enumerate(RELEASE_HOURS):
        s = int(rh * 6)
        for d in range(DAYS):
            g = pv_forecast_10min(corrected, pv, d, k)
            # 当前版本残差使用原日前负荷预测减更新光伏，未包含负荷反馈修正
            out[k, d, s:] = (load[d, s:] - pv[d, s:]) - (pred_load[d, s:] - g)
    return out


def reserve_version(version_residuals: np.ndarray, day: int, version: int):
    s = int(RELEASE_HOURS[version] * 6)
    if day == 0:
        return np.zeros(N_SLOT - s)
    # 同一版本只使用此前日期的剩余时段残差，当前分位数固定为0.5
    hist = version_residuals[version, max(0, day-28):day, s:]
    return np.maximum(np.quantile(hist, 0.5, axis=0), 0.0)


def load_feedback(pred_load: np.ndarray, load: np.ndarray, day: int, start: int):
    if start == 0:
        return pred_load[day].copy()
    # 用已发生的过去12时段负荷偏差修正未来，36时段对应六小时衰减尺度
    bias = float(np.mean(load[day, start-12:start] - pred_load[day, start-12:start]))
    t = np.arange(start, N_SLOT)
    return np.maximum(pred_load[day, start:] + bias * np.exp(-(t-start)/36.0), 0.0)


def price_feedback(pred_price: np.ndarray, actual_price: np.ndarray, day: int, start: int):
    if start == 0:
        return pred_price[day].copy()
    # 只用已发生的过去两小时价格偏差，未来真实价格不进入反馈计算
    bias = float(np.mean(actual_price[day, start-12:start] - pred_price[day, start-12:start]))
    t = np.arange(start, N_SLOT)
    return np.maximum(pred_price[day, start:] + bias * np.exp(-(t-start)/36.0), 0.001)


# ============================================================================
# result 模板写入
# ============================================================================
def positive_runs(arr: np.ndarray, eps: float = 1e-7):
    runs = []
    i = 0
    while i < len(arr):
        if arr[i] > eps:
            j = i
            total = 0.0
            while j < len(arr) and arr[j] > eps:
                total += float(arr[j])
                j += 1
            # 区间采用左闭右开下标，total为所有连续应急时段电量之和
            runs.append((i, j, total))
            i = j
        else:
            i += 1
    return runs


def _clock(slot_boundary: int):
    """10分钟边界序号转为 HH:MM；144 显示为 24:00。"""
    m = slot_boundary * 10
    if m == 24 * 60:
        return "24:00"
    return f"{m // 60}:{m % 60:02d}"


def _excel_date_to_day_index(value: Any, epoch=None):
    """把模板日期单元格转换为相对 2025-01-01 的日序号；非日期返回 None。"""
    if value is None or value == "⁝":
        return None
    if isinstance(value, datetime):
        return (value.date() - BASE_DATE).days
    if isinstance(value, date):
        return (value - BASE_DATE).days
    # Excel日期可能是序列数，按工作簿epoch转换以兼容日期基准
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            d = from_excel(value, epoch=epoch)
            if isinstance(d, datetime):
                d = d.date()
            return (d - BASE_DATE).days
        except Exception:
            # 日期转换失败时沿用固定序列数基准估计日序号
            return int(round(float(value) - 45658.0))
    if isinstance(value, str):
        text = value.strip().replace(".", "-").replace("/", "-")
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return (datetime.strptime(text, fmt).date() - BASE_DATE).days
            except ValueError:
                pass
    return None


def _fill_charge_sheet(wb, ch: np.ndarray, dis: np.ndarray, E: np.ndarray):
    ws = wb["充放电量"]
    blocks = {
        "0:00-4:00": 0,
        "4:00-8:00": 24,
        "8:00-12:00": 48,
        "12:00-16:00": 72,
        "16:00-20:00": 96,
        "20:00-24:00": 120,
    }
    current_day = None
    # 按模板现有行扫描；只填模板已经列出的指定日期，保留原格式。
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == "⁝":
            current_day = None
        else:
            idx = _excel_date_to_day_index(a, wb.epoch)
            if idx is not None:
                current_day = idx
        if current_day is None or not (0 <= current_day < ch.shape[0]):
            continue

        period = ws.cell(r, 2).value
        if period in blocks:
            s = blocks[period]
            ws.cell(r, 3).value = float(ch[current_day, s:s+24].sum())
            ws.cell(r, 4).value = float(dis[current_day, s:s+24].sum())

        marker = ws.cell(r, 5).value
        # 当前分支仅按下面列出的零点和24时形式匹配时刻标记
        marker_text = str(marker).strip() if marker is not None else ""
        if marker == 0 or marker_text == "0:00":
            ws.cell(r, 6).value = float(E[current_day, 0])
        elif marker_text == "24:00":
            ws.cell(r, 6).value = float(E[current_day, -1])


def _fill_emergency_sheet(wb, emergency: np.ndarray):
    ws = wb["紧急购电量"]
    r = 2
    max_row = ws.max_row
    while r <= max_row:
        day = _excel_date_to_day_index(ws.cell(r, 1).value, wb.epoch)
        if day is None or not (0 <= day < emergency.shape[0]):
            r += 1
            continue

        # 当前日期在模板中占用到下一条日期之前的所有空白行。
        rr = r + 1
        while rr <= max_row and ws.cell(rr, 1).value is None:
            rr += 1
        # 预留行数限制当前日期最多能写入多少段应急记录
        capacity = rr - r

        # 先清理模板中的旧示例值，避免残留。
        for row in range(r, r + capacity):
            ws.cell(row, 2).value = None
            ws.cell(row, 3).value = None

        runs = positive_runs(emergency[day])
        # 当前实现只填入预留行能容纳的前几个区间，其余不会写入模板
        for j, (s, e, amount) in enumerate(runs[:capacity]):
            ws.cell(r + j, 2).value = f"{_clock(s)}-{_clock(e)}"
            ws.cell(r + j, 3).value = float(amount)
        r = rr


def _write_matrix(ws, start_row: int, start_col: int, matrix):
    for i, row in enumerate(matrix, start=start_row):
        for j, value in enumerate(row, start=start_col):
            ws.cell(i, j).value = value


def _template_purchase_values(q, day=None, next_day_first=None):
    # 模板覆盖当天00:10至次日00:10，因此依次填入当天索引1至143和次日索引0
    values = np.asarray(q, dtype=float)
    if day is None:
        if values.shape != (144,):
            raise ValueError("典型日购电数组必须包含144个十分钟时段")
        # 问题1每天条件相同，因此次日首时段复用典型日首时段
        return values[1:].tolist() + [float(values[0])]
    if values.ndim != 2 or values.shape[1] != 144:
        raise ValueError("逐日购电数组必须采用日期乘144时段的二维形状")
    if day + 1 < len(values):
        tail = float(values[day + 1, 0])
    elif next_day_first is not None:
        tail = float(next_day_first)
        if not np.isfinite(tail) or tail < 0:
            raise ValueError("补充的次日首时段购电量必须是非负有限数值")
    else:
        import warnings
        # 缺失次年首时段时保留空白，不能以零或本年首日数据替代未知购电量
        warnings.warn("缺少最后一天的次日00:00—00:10购电量，模板对应末格留空", RuntimeWarning, stacklevel=2)
        tail = None
    return values[day, 1:].tolist() + [tail]


def write_result1(template_path: str | Path, output_path: str | Path,
                  q, ch, dis, E):
    wb = load_workbook(template_path)
    ws_q = wb["计划购电量"]
    # 保留模板标签，按典型日跨日顺序重新定位购电值
    for i, value in enumerate(_template_purchase_values(q), start=2):
        ws_q.cell(i, 2).value = float(value)

    ws_cd = wb["充放电量"]
    # 每组24个十分钟时段对应四小时，按六组汇总充放电电量
    for b in range(6):
        s = 24 * b
        ws_cd.cell(2 + b, 2).value = float(ch[s:s+24].sum())
        ws_cd.cell(2 + b, 3).value = float(dis[s:s+24].sum())
    ws_cd.cell(2, 5).value = float(E[0])
    ws_cd.cell(3, 5).value = float(E[-1])

    wb.save(output_path)
    wb.close()


def write_result2(template_path: str | Path, output_path: str | Path,
                  q, ch, dis, emergency, E, daily_cost, *, next_day_first=None):
    wb = load_workbook(template_path)
    ws = wb["计划购电量"]
    # 时段列按模板跨日映射，日总量和费用仍按原来的日历日口径统计
    rows = [
        _template_purchase_values(q, d, next_day_first) + [float(q[d].sum()), float(daily_cost[d])]
        for d in range(31, 365)
    ]
    _write_matrix(ws, 2, 2, rows)
    _fill_charge_sheet(wb, ch, dis, E)
    _fill_emergency_sheet(wb, emergency)
    wb.save(output_path)
    wb.close()


def write_result3(template_path: str | Path, output_path: str | Path,
                  q0, q, ch, dis, emergency, E, daily_cost, settlement_price,
                  *, next_day_first_plan=None, next_day_first_adjusted=None):
    wb = load_workbook(template_path)
    plan_rows = []
    adjust_rows = []
    for d in range(31, 365):
        # 原计划表末列只计算原日前购电费，不包含日内调整和应急费用
        plan_rows.append(
            _template_purchase_values(q0, d, next_day_first_plan)
            + [float(q0[d].sum()), float((settlement_price[d] * q0[d]).sum())]
        )
        # 调整表写最终常规购电量并在末列写入包含应急的实际总费用
        adjust_rows.append(
            _template_purchase_values(q, d, next_day_first_adjusted)
            + [float(q[d].sum()), float(daily_cost[d])]
        )

    _write_matrix(wb["计划购电量"], 2, 2, plan_rows)
    _write_matrix(wb["调整购电量"], 2, 2, adjust_rows)
    _fill_charge_sheet(wb, ch, dis, E)
    _fill_emergency_sheet(wb, emergency)
    wb.save(output_path)
    wb.close()
